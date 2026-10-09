import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Avg, Q
from django.contrib.auth import get_user_model

from apps.businesses.models import Negocio
from apps.billing.models import Suscripcion, Plan, EstadoSuscripcionChoices, CambioPlan
from apps.geo.models import Rubro
from apps.encuestas.models import RespuestaEncuesta, EmocionCSATChoices
from apps.reputation.models import ResenaGoogle, SentimientoChoices

User = get_user_model()


def humanize_antiguedad(fecha):
    if not fecha:
        return "N/I"
    now = timezone.now()
    if timezone.is_naive(fecha):
        fecha = timezone.make_aware(fecha)
    diff = now - fecha
    days = diff.days
    if days < 1:
        return "Hoy"
    if days < 30:
        return f"{days} día{'s' if days > 1 else ''}"
    months = days // 30
    years = months // 12
    rem_months = months % 12
    if years > 0:
        if rem_months > 0:
            return f"{years} año{'s' if years > 1 else ''} {rem_months} mes{'es' if rem_months > 1 else ''}"
        return f"{years} año{'s' if years > 1 else ''}"
    return f"{months} mes{'es' if months > 1 else ''}"


def _crear_workbook_estilizado(titulo_reporte):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = titulo_reporte[:31]

    # Estilos de marca ClientBeat
    fill_header = PatternFill(start_color="0B3EA6", end_color="0B3EA6", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_title = Font(name="Calibri", size=16, bold=True, color="0B3EA6")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="64748B")
    border_thin = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Título principal del reporte
    ws.merge_cells("A1:G1")
    ws["A1"] = f"ClientBeat · {titulo_reporte}"
    ws["A1"].font = font_title

    # Subtítulo con fecha
    now_str = timezone.localtime(timezone.now()).strftime("%d/%m/%Y %H:%M:%S")
    ws.merge_cells("A2:G2")
    ws["A2"] = f"Reporte generado automáticamente el {now_str} (Chile)"
    ws["A2"].font = font_subtitle

    return wb, ws, fill_header, font_header, border_thin


def _ajustar_y_responder(wb, filename):
    for ws in wb.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if cell.row < 4:  # Ignorar título largo
                    continue
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    from io import BytesIO
    buffer = BytesIO()
    wb.save(buffer)
    file_data = buffer.getvalue()

    response = HttpResponse(
        file_data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Content-Length'] = len(file_data)
    return response


# ==========================================
# EXPORT 1: CLIENTES Y ANTIGÜEDAD
# ==========================================
def export_clientes_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Reporte de Clientes y Antigüedad")

    headers = [
        "Negocio", "RUT", "Dueño (Email)", "Rubro", "Comuna", "Locales",
        "Fecha Registro", "Antigüedad Negocio", "Estado Negocio",
        "Plan Activo", "Estado Suscripción", "Fecha Inicio Plan", "Antigüedad en Plan"
    ]

    ws.append([])  # Fila 3 vacía
    ws.append(headers)  # Fila 4 headers

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    negocios = (
        Negocio.objects.select_related('dueño', 'comuna', 'rubro')
        .annotate(total_locales=Count('locales', distinct=True))
        .prefetch_related('suscripciones__plan')
        .order_by('-fecha_creacion')
    )

    q = request.GET.get('q', '').strip()
    if q:
        negocios = negocios.filter(Q(nombre__icontains=q) | Q(rut__icontains=q) | Q(dueño__email__icontains=q))

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, n in enumerate(negocios, start=5):
        sus_activa = n.suscripciones.filter(estado=EstadoSuscripcionChoices.ACTIVA).first()
        plan_nombre = sus_activa.plan.get_nombre_mostrar() if sus_activa and sus_activa.plan else "Sin plan activo"
        sus_estado = sus_activa.get_estado_display() if sus_activa else "Inactiva"
        fecha_inicio_plan = sus_activa.fecha_inicio.strftime("%d/%m/%Y") if sus_activa and sus_activa.fecha_inicio else "-"
        antiguedad_plan = humanize_antiguedad(sus_activa.fecha_inicio) if sus_activa and sus_activa.fecha_inicio else "-"

        row = [
            n.nombre,
            n.rut or "-",
            n.dueño.email if n.dueño else "Sin asignar",
            n.rubro.nombre if n.rubro else "Sin rubro",
            n.comuna.nombre if n.comuna else "-",
            n.total_locales,
            n.fecha_creacion.strftime("%d/%m/%Y %H:%M"),
            humanize_antiguedad(n.fecha_creacion),
            n.get_estado_display(),
            plan_nombre,
            sus_estado,
            fecha_inicio_plan,
            antiguedad_plan
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Clientes_Antiguedad_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)


# ==========================================
# EXPORT 2: BENCHMARK GOOGLE POR RUBRO
# ==========================================
def export_benchmark_google_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Benchmark Reseñas Google por Rubro")

    headers = [
        "Rubro", "Categoría Google", "Negocios Registrados", "Locales Evaluados",
        "Total Reseñas Google", "Rating Promedio Google (1-5★)", "Sentimiento Positivo (%)",
        "Sentimiento Neutro (%)", "Sentimiento Negativo (%)"
    ]

    ws.append([])
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rubros = Rubro.objects.filter(activo=True).annotate(
        total_negocios=Count('negocios', distinct=True)
    ).order_by('-total_negocios')

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, r in enumerate(rubros, start=5):
        resenas = ResenaGoogle.objects.filter(local__negocio__rubro=r)
        total_res = resenas.count()
        avg_rating = resenas.aggregate(avg=Avg('calificacion'))['avg'] or 0.0
        locales_cnt = n_locales = Negocio.objects.filter(rubro=r).aggregate(locs=Count('locales'))['locs'] or 0

        pos = resenas.filter(sentimiento=SentimientoChoices.POSITIVO).count()
        neu = resenas.filter(sentimiento=SentimientoChoices.NEUTRO).count()
        neg = resenas.filter(sentimiento=SentimientoChoices.NEGATIVO).count()

        pos_pct = f"{round(100*pos/total_res, 1)}%" if total_res else "0.0%"
        neu_pct = f"{round(100*neu/total_res, 1)}%" if total_res else "0.0%"
        neg_pct = f"{round(100*neg/total_res, 1)}%" if total_res else "0.0%"

        cat_g = r.google_category.google_es if hasattr(r, 'google_category') and r.google_category else "-"

        row = [
            r.nombre,
            cat_g,
            r.total_negocios,
            locales_cnt,
            total_res,
            round(avg_rating, 2),
            pos_pct,
            neu_pct,
            neg_pct,
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Benchmark_Google_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)


# ==========================================
# EXPORT 3: BENCHMARK CLIENTBEAT POR RUBRO
# ==========================================
def export_benchmark_clientbeat_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Benchmark ClientBeat Encuestas QR")

    headers = [
        "Rubro", "Negocios", "Respuestas QR Totales", "CSAT % Felices",
        "NPS Net Score (-100 a +100)", "Promotores (9-10)", "Pasivos (7-8)",
        "Detractores (0-6)", "Clasificación Nivel"
    ]

    ws.append([])
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rubros = Rubro.objects.filter(activo=True).annotate(
        total_negocios=Count('negocios', distinct=True)
    ).order_by('-total_negocios')

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, r in enumerate(rubros, start=5):
        respuestas = RespuestaEncuesta.objects.filter(local__negocio__rubro=r)
        total_resp = respuestas.count()

        # CSAT
        felices = respuestas.filter(csat_emocion__in=['MUY_FELIZ', 'FELIZ', 'MUY_SATISFECHO']).count()
        csat_pct = f"{round(100*felices/total_resp, 1)}%" if total_resp else "0.0%"

        # NPS
        resp_nps = respuestas.filter(nps_puntaje__isnull=False)
        total_nps = resp_nps.count()
        if total_nps:
            p = resp_nps.filter(nps_puntaje__gte=9).count()
            pas = resp_nps.filter(nps_puntaje__in=[7, 8]).count()
            d = resp_nps.filter(nps_puntaje__lte=6).count()
            score = round(100 * (p - d) / total_nps)
            clasificacion = "Sobresaliente (≥50)" if score >= 50 else ("Aceptable (0-49)" if score >= 0 else "En Riesgo (<0)")
        else:
            p = pas = d = score = 0
            clasificacion = "Sin datos"

        row = [
            r.nombre,
            r.total_negocios,
            total_resp,
            csat_pct,
            score,
            p,
            pas,
            d,
            clasificacion
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Benchmark_ClientBeat_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)


# ==========================================
# EXPORT 4: BENCHMARK NPS & CSAT DEDICADO
# ==========================================
def export_benchmark_nps_csat_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Reporte Benchmark NPS y CSAT por Rubro")

    headers = [
        "Rubro", "Empresas", "Evaluaciones Totales", "CSAT Felices %",
        "CSAT Neutral %", "CSAT Insatisfecho %", "NPS Net Score",
        "% Promotores", "% Pasivos", "% Detractores", "Nivel Excelencia"
    ]

    ws.append([])
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rubros = Rubro.objects.filter(activo=True).annotate(
        total_negocios=Count('negocios', distinct=True)
    ).order_by('-total_negocios')

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, r in enumerate(rubros, start=5):
        resp = RespuestaEncuesta.objects.filter(local__negocio__rubro=r)
        total_resp = resp.count()

        felices = resp.filter(csat_emocion__in=['MUY_FELIZ', 'FELIZ', 'MUY_SATISFECHO']).count()
        neutrales = resp.filter(csat_emocion__in=['NEUTRAL']).count()
        insatisfechos = resp.filter(csat_emocion__in=['INSATISFECHO', 'MUY_INSATISFECHO']).count()

        csat_fel = f"{round(100*felices/total_resp, 1)}%" if total_resp else "0.0%"
        csat_neu = f"{round(100*neutrales/total_resp, 1)}%" if total_resp else "0.0%"
        csat_ins = f"{round(100*insatisfechos/total_resp, 1)}%" if total_resp else "0.0%"

        resp_nps = resp.filter(nps_puntaje__isnull=False)
        total_nps = resp_nps.count()
        if total_nps:
            p = resp_nps.filter(nps_puntaje__gte=9).count()
            pas = resp_nps.filter(nps_puntaje__in=[7, 8]).count()
            d = resp_nps.filter(nps_puntaje__lte=6).count()
            nps_score = round(100 * (p - d) / total_nps)
            p_pct = f"{round(100*p/total_nps, 1)}%"
            pas_pct = f"{round(100*pas/total_nps, 1)}%"
            d_pct = f"{round(100*d/total_nps, 1)}%"
            nivel = "Excelente (≥50)" if nps_score >= 50 else ("Bueno (0-49)" if nps_score >= 0 else "En Riesgo (<0)")
        else:
            nps_score = 0
            p_pct = pas_pct = d_pct = "0.0%"
            nivel = "Sin evaluaciones"

        row = [
            r.nombre,
            r.total_negocios,
            total_resp,
            csat_fel,
            csat_neu,
            csat_ins,
            nps_score,
            p_pct,
            pas_pct,
            d_pct,
            nivel
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Benchmark_NPS_CSAT_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)


# ==========================================
# EXPORT 5: TENDENCIAS EN RESEÑAS
# ==========================================
def export_tendencias_resenas_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Reporte de Tendencias en Reseñas")

    headers = [
        "Negocio", "Local", "Autor", "Calificación (1-5★)", "Sentimiento",
        "Fecha Reseña Google", "Comentario"
    ]

    ws.append([])
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    resenas = ResenaGoogle.objects.select_related('local', 'local__negocio').order_by('-fecha_google')[:1000]

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, r in enumerate(resenas, start=5):
        row = [
            r.local.negocio.nombre if r.local and r.local.negocio else "-",
            r.local.nombre if r.local else "-",
            r.autor_nombre,
            r.calificacion,
            r.get_sentimiento_display() if r.sentimiento else "Sin clasificar",
            r.fecha_google.strftime("%d/%m/%Y %H:%M") if r.fecha_google else "-",
            r.comentario or "-"
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Tendencias_Resenas_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)


# ==========================================
# EXPORT 6: PLANES Y SUSCRIPCIONES
# ==========================================
def export_planes_excel(request):
    wb, ws, fill_header, font_header, border_thin = _crear_workbook_estilizado("Reporte de Planes y Suscripciones")

    headers = [
        "Negocio", "RUT", "Plan", "Estado Suscripción", "Precio Base CLP",
        "Monto Estimado CLP", "Fecha Inicio", "Fecha Vencimiento", "Antigüedad Suscripción"
    ]

    ws.append([])
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    suscripciones = Suscripcion.objects.select_related('negocio', 'plan').order_by('-fecha_inicio')

    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    for i, s in enumerate(suscripciones, start=5):
        monto = s.plan.precio_clp if (s.plan and s.plan.precio_clp is not None) else 0
        venc = s.fecha_vencimiento.strftime("%d/%m/%Y") if s.fecha_vencimiento else "Sin fecha"
        antiguedad = humanize_antiguedad(s.fecha_inicio) if s.fecha_inicio else "-"

        row = [
            s.negocio.nombre if s.negocio else "-",
            s.negocio.rut if s.negocio else "-",
            s.plan.get_nombre_mostrar() if s.plan else "Sin plan",
            s.get_estado_display(),
            s.plan.precio_clp if s.plan else 0,
            monto,
            s.fecha_inicio.strftime("%d/%m/%Y") if s.fecha_inicio else "-",
            venc,
            antiguedad
        ]
        ws.append(row)

        for col_num in range(1, len(row) + 1):
            cell = ws.cell(row=i, column=col_num)
            cell.border = border_thin
            if i % 2 == 0:
                cell.fill = fill_zebra

    filename = f"ClientBeat_Planes_Suscripciones_{timezone.now().strftime('%Y%m%d')}.xlsx"
    return _ajustar_y_responder(wb, filename)
